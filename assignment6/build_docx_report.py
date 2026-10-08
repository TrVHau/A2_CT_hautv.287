"""
Script tạo báo cáo Word hoàn chỉnh 20+ trang cho Assignment 6
Sinh viên: Trần Văn Hậu - B23DCCN287
Giảng viên: Trần Đình Quế
"""
import os
os.environ["TF_CPP_MIN_LOG_LEVEL"] = "3"
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

ROOT = Path(__file__).resolve().parent
FIG = ROOT / "figures"
OUT = ROOT / "Bao_Cao_Assignment_6_Tran_Van_Hau_B23DCCN287.docx"

COLOR_PRIMARY = RGBColor(27, 54, 93)
COLOR_SECONDARY = RGBColor(41, 128, 185)
COLOR_TEXT = RGBColor(33, 37, 41)
COLOR_MUTED = RGBColor(108, 117, 125)

def set_cell_shading(cell, hex_color):
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    cell._tc.get_or_add_tcPr().append(shd)

def add_table(doc, headers, data, widths=None):
    t = doc.add_table(rows=len(data)+1, cols=len(headers))
    for i, h in enumerate(headers):
        c = t.rows[0].cells[i]
        c.text = h
        set_cell_shading(c, "1B365D")
        for r in c.paragraphs[0].runs:
            r.font.bold, r.font.color.rgb, r.font.size = True, RGBColor(255,255,255), Pt(10)
    for ri, row in enumerate(data):
        for ci, val in enumerate(row):
            t.rows[ri+1].cells[ci].text = str(val)
            set_cell_shading(t.rows[ri+1].cells[ci], "F4F6F9" if ri%2 else "FFFFFF")
    doc.add_paragraph()
    return t

def add_fig(doc, path, caption):
    if Path(path).exists():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(path), width=Inches(5.8))
        pc = doc.add_paragraph(caption)
        pc.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for r in pc.runs:
            r.font.italic, r.font.bold, r.font.size, r.font.color.rgb = True, True, Pt(9.5), COLOR_PRIMARY

def add_h1(doc, text):
    h = doc.add_heading(text, 1)
    for r in h.runs:
        r.font.name, r.font.size, r.font.color.rgb = "Times New Roman", Pt(15), COLOR_PRIMARY

def add_h2(doc, text):
    h = doc.add_heading(text, 2)
    for r in h.runs:
        r.font.name, r.font.size, r.font.color.rgb = "Times New Roman", Pt(13), COLOR_SECONDARY

def add_p(doc, text):
    p = doc.add_paragraph(text)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    for r in p.runs:
        r.font.name, r.font.size, r.font.color.rgb = "Times New Roman", Pt(11.5), COLOR_TEXT

doc = Document()
for s in doc.sections:
    s.page_width, s.page_height = Inches(8.27), Inches(11.69)
    s.top_margin, s.bottom_margin, s.left_margin, s.right_margin = Inches(0.8), Inches(0.8), Inches(1.0), Inches(0.8)

# TRANG BÌA
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("HỌC VIỆN CÔNG NGHỆ BƯU CHÍNH VIỄN THÔNG\nKHOA CÔNG NGHỆ THÔNG TIN\n" + "-"*40)
r.font.bold, r.font.size = True, Pt(13)
r.font.color.rgb = COLOR_PRIMARY

doc.add_paragraph()
pt = doc.add_paragraph()
pt.alignment = WD_ALIGN_PARAGRAPH.CENTER
rt = pt.add_run("BÁO CÁO ASSIGNMENT 6\n\n")
rt.font.bold, rt.font.size, rt.font.color.rgb = True, Pt(14), COLOR_SECONDARY
rt2 = pt.add_run("DỰ BÁO CHUỖI THỜI GIAN VỚI MẠNG NƠ-RON HỒI QUY\n(LSTM - PYTORCH vs KERAS)")
rt2.font.bold, rt2.font.size, rt2.font.color.rgb = True, Pt(17), COLOR_PRIMARY

doc.add_paragraph()
info_headers = ["THÔNG TIN", "CHI TIẾT"]
info_data = [
    ["Sinh viên", "Trần Văn Hậu"],
    ["Mã số sinh viên", "B23DCCN287"],
    ["Giảng viên", "Trần Đình Quế"],
    ["Học phần", "Học máy & Học sâu"],
    ["Dataset Chứng khoán", "https://www.kaggle.com/datasets/camnugent/sandp500"],
    ["Dataset Bán lẻ", "https://www.kaggle.com/datasets/mashlyn/online-retail-ii-uci"],
    ["GitHub Repository", "https://github.com/TrVHau/A2_CT_hautv.287.git"],
    ["Ngày hoàn thành", "30/09/2026"]
]
add_table(doc, info_headers, info_data)

doc.add_page_break()

# TÓM TẮT
add_h1(doc, "TÓM TẮT")
add_p(doc, "Báo cáo trình bày đầy đủ quy trình xây dựng mô hình LSTM dự báo chuỗi thời gian trên hai bộ dữ liệu thực tế: giá cổ phiếu Apple (AAPL) từ S&P 500 và giao dịch Online Retail II. Mỗi framework (PyTorch và Keras) dùng kiến trúc LSTM(32) -> Dense(16, ReLU) -> Dense(1) với cửa sổ 30 bước.")
add_p(doc, "Kết quả thực nghiệm từ notebook cho thấy Keras đạt MAE 6.88/RMSE 8.16 trên AAPL và MAE 15.32/RMSE 19.82 trên Retail. PyTorch đạt MAE 47.46/RMSE 48.51 và MAE 23.80/RMSE 29.84 tương ứng. PyTorch nhanh hơn 3.4 lần trên Stock và 7.7 lần trên Transactions. Cả hai mô hình được triển khai thành FastAPI tại port 8001 (PyTorch) và 8002 (Keras) với giao diện web tương tác.")

doc.add_page_break()

# CHƯƠNG 1
add_h1(doc, "1. MỤC TIÊU VÀ PHẠM VI")
add_h2(doc, "1.1. Yêu cầu đề bài")
add_p(doc, "Đề bài yêu cầu 5 nội dung: (1) Trình bày lý thuyết RNN, biểu diễn hàm và code; (2) Chọn hai tập dữ liệu (chứng khoán và giao dịch mua hàng), mô tả và phân bố; (3) Code RNN bằng PyTorch và deploy web; (4) Code RNN bằng Keras và deploy web; (5) So sánh hai mô hình.")

add_h2(doc, "1.2. Nguồn dữ liệu")
add_p(doc, "• Chứng khoán S&P 500 (AAPL): https://www.kaggle.com/datasets/camnugent/sandp500 - 1.259 phiên từ 2013-02-08 đến 2018-02-07.")
add_p(doc, "• Online Retail II: https://www.kaggle.com/datasets/mashlyn/online-retail-ii-uci - 604 ngày từ 2009-12-01 đến 2011-12-09 sau làm sạch.")
add_p(doc, "• GitHub: https://github.com/TrVHau/A2_CT_hautv.287.git")

doc.add_page_break()

# CHƯƠNG 2
add_h1(doc, "2. LÝ THUYẾT RNN VÀ LSTM")
add_h2(doc, "2.1. Chuỗi thời gian")
add_p(doc, "Chuỗi thời gian là dãy quan sát theo thứ tự: y_1, y_2, ..., y_T. Khác với dữ liệu bảng, thứ tự không được xáo trộn. Với cửa sổ L=30, mô hình nhận X_i=[y_i,...,y_{i+29}] và dự đoán y_{i+30}.")

add_h2(doc, "2.2. Vanilla RNN")
add_p(doc, "RNN cập nhật trạng thái ẩn: h_t = tanh(W_x*x_t + W_h*h_{t-1} + b_h). Vấn đề: biến mất/bùng nổ đạo hàm khi lan truyền ngược qua nhiều bước thời gian.")

add_h2(doc, "2.3. LSTM")
add_p(doc, "LSTM có 2 trạng thái (h_t, c_t) và 3 cổng điều khiển:")
add_p(doc, "• Forget gate: f_t = σ(W_f*[h_{t-1}, x_t] + b_f)")
add_p(doc, "• Input gate: i_t = σ(W_i*[h_{t-1}, x_t] + b_i), c_tilde = tanh(W_c*[h_{t-1}, x_t] + b_c)")
add_p(doc, "• Cell update: c_t = f_t ⊙ c_{t-1} + i_t ⊙ c_tilde")
add_p(doc, "• Output gate: o_t = σ(W_o*[h_{t-1}, x_t] + b_o), h_t = o_t ⊙ tanh(c_t)")
add_p(doc, "Cell state c_t tạo đường truyền thẳng giúp gradient không bị suy giảm.")

add_fig(doc, FIG/"rnn_vs_lstm_cell.png", "Hình 2.1: So sánh Vanilla RNN và LSTM Cell")
add_fig(doc, FIG/"lstm_architecture_diagram.png", "Hình 2.2: Kiến trúc LSTM Regressor Many-to-One")

doc.add_page_break()

# CHƯƠNG 3
add_h1(doc, "3. DỮ LIỆU VÀ EDA")
add_h2(doc, "3.1. Dữ liệu AAPL")
add_p(doc, "Lọc mã AAPL từ S&P 500, sắp xếp theo date, giữ close làm target. Tổng 1.259 phiên, 80% đầu (1.007) train, 20% cuối (252) test.")

stat_h = ["Chỉ số", "Giá trị"]
stat_d = [
    ["Mean", "109.0667 USD"],
    ["Std", "30.5568"],
    ["Min", "55.7899"],
    ["Q1", "84.8306"],
    ["Median", "109.0100"],
    ["Q3", "127.1200"],
    ["Max", "179.2600"],
    ["Skewness", "0.2899"],
    ["Kurtosis", "-0.6093"]
]
add_table(doc, stat_h, stat_d)

add_fig(doc, FIG/"aapl_time_series.png", "Hình 3.1: Chuỗi giá AAPL và ranh giới Train/Test")
add_fig(doc, FIG/"aapl_distribution.png", "Hình 3.2: Phân bố AAPL (Histogram, KDE, Boxplot)")

add_h2(doc, "3.2. Dữ liệu Online Retail II")
add_p(doc, "Từ 1.067.371 dòng thô, loại invoice bắt đầu 'C', Quantity<=0, Price<=0, dropna. Tổng hợp số invoice duy nhất/ngày, còn 604 ngày.")

ret_h = ["Chỉ số", "Invoice/ngày"]
ret_d = [
    ["Mean", "66.3526"],
    ["Std", "24.7778"],
    ["Min", "11"],
    ["Median", "63"],
    ["Max", "153"],
    ["Skewness", "0.6803"],
    ["Kurtosis", "0.3740"]
]
add_table(doc, ret_h, ret_d)

add_fig(doc, FIG/"retail_time_series.png", "Hình 3.3: Chuỗi Invoice/ngày và ranh giới Train/Test")
add_fig(doc, FIG/"retail_distribution.png", "Hình 3.4: Phân bố Retail (Histogram, KDE, Boxplot)")

doc.add_page_break()

# CHƯƠNG 4
add_h1(doc, "4. TIỀN XỬ LÝ")
add_h2(doc, "4.1. Chia dữ liệu theo thời gian")
add_p(doc, "80% đầu train, 20% cuối test, không shuffle để giữ thứ tự. AAPL: 1.007 train, 252 test. Retail: 483 train, 121 test.")

add_h2(doc, "4.2. Chuẩn hóa MinMaxScaler")
add_p(doc, "Scaler chỉ fit trên train để tránh leakage. Công thức: x_scaled = (x - min_train) / (max_train - min_train). Sau đó transform toàn bộ chuỗi.")

add_h2(doc, "4.3. Cửa sổ trượt")
add_p(doc, "Với window=30, mỗi mẫu X có shape (30, 1) và y có shape (1). Số mẫu train thực tế: 977 (AAPL), 453 (Retail).")

add_h2(doc, "4.4. Inverse transform")
add_p(doc, "Sau dự đoán, cả y_test và prediction được inverse về đơn vị gốc trước khi tính MAE/RMSE/MAPE.")

doc.add_page_break()

# CHƯƠNG 5
add_h1(doc, "5. MÔ HÌNH PYTORCH")
add_h2(doc, "5.1. Kiến trúc")
add_p(doc, "Lớp LSTMRegressor trong train_pytorch.py:")
add_p(doc, "• LSTM(input_size=1, hidden_size=32, batch_first=True)")
add_p(doc, "• Head: Linear(32, 16) -> ReLU -> Linear(16, 1)")
add_p(doc, "• Forward: lấy output[:, -1, :] rồi đưa qua head")

add_h2(doc, "5.2. Huấn luyện")
add_p(doc, "Optimizer Adam(lr=0.003), loss MSE, 18 epochs. Seed 42 cho torch và random. Model.eval() + torch.no_grad() khi đánh giá.")

add_h2(doc, "5.3. Artifact")
add_p(doc, "Lưu state_dict vào .pth và scaler JSON. models/pytorch_stock.pth, models/pytorch_stock_scaler.json, tương tự cho transactions.")

doc.add_page_break()

# CHƯƠNG 6
add_h1(doc, "6. MÔ HÌNH KERAS")
add_h2(doc, "6.1. Kiến trúc")
add_p(doc, "Sequential trong train_keras.py:")
add_p(doc, "• Input(shape=(30, 1))")
add_p(doc, "• LSTM(32, return_sequences=False)")
add_p(doc, "• Dense(16, activation='relu')")
add_p(doc, "• Dense(1, activation='linear')")
add_p(doc, "• Compile: Adam(lr=0.003), loss='mse'")

add_h2(doc, "6.2. Huấn luyện")
add_p(doc, "tf.keras.utils.set_random_seed(42). Fit với batch_size=64, epochs=18, verbose=0. Predict với verbose=0.")

add_h2(doc, "6.3. Artifact")
add_p(doc, "Lưu model.save() thành .keras và scaler JSON. models/keras_stock.keras, models/keras_stock_scaler.json, tương tự cho transactions.")

doc.add_page_break()

# CHƯƠNG 7
add_h1(doc, "7. KẾT QUẢ VÀ SO SÁNH")
add_h2(doc, "7.1. Bảng kết quả notebook")

comp_h = ["Framework", "Dataset", "Seconds", "MAE", "RMSE", "MAPE (%)"]
comp_d = [
    ["Keras", "stock", "12.61", "6.8849", "8.1645", "4.2579"],
    ["PyTorch", "stock", "3.71", "47.4643", "48.5066", "30.1820"],
    ["Keras", "transactions", "8.58", "15.3181", "19.8245", "24.5338"],
    ["PyTorch", "transactions", "1.11", "23.7989", "29.8393", "31.1915"]
]
add_table(doc, comp_h, comp_d)

add_h2(doc, "7.2. Phân tích")
add_p(doc, "Keras tốt hơn về MAE/RMSE do dùng mini-batch SGD (batch=64), giúp cập nhật nhiều lần/epoch và hội tụ tốt hơn. PyTorch dùng full-batch nên chỉ cập nhật 18 lần trong 18 epochs, chưa hội tụ đủ.")
add_p(doc, "PyTorch nhanh hơn 3.4x (stock) và 7.7x (transactions) do ít overhead và không chia batch.")

add_fig(doc, FIG/"stock_predictions_comparison.png", "Hình 7.1: So sánh dự báo AAPL")
add_fig(doc, FIG/"retail_predictions_comparison.png", "Hình 7.2: So sánh dự báo Retail")
add_fig(doc, FIG/"metrics_comparison_bar.png", "Hình 7.3: So sánh MAE, RMSE, Seconds")

doc.add_page_break()

# CHƯƠNG 8
add_h1(doc, "8. DEPLOYMENT WEB")
add_h2(doc, "8.1. Kiến trúc")
add_p(doc, "Hai FastAPI service:")
add_p(doc, "• web/pytorch_app.py: Port 8001, nạp .pth + scaler JSON")
add_p(doc, "• web/keras_app.py: Port 8002, nạp .keras + scaler JSON")
add_p(doc, "• web/index.html: Giao diện cho cả hai")

add_h2(doc, "8.2. Endpoint")
add_p(doc, "• GET /: Trả index.html")
add_p(doc, "• GET /health: Trạng thái và danh sách model")
add_p(doc, "• GET /series/{dataset}: 30 giá trị cuối")
add_p(doc, "• POST /predict: Nhận dataset và values (list 30 số hoặc rỗng), trả prediction")
add_p(doc, "• GET /docs: Swagger UI")

add_h2(doc, "8.3. Chạy")
add_p(doc, "uvicorn web.pytorch_app:app --app-dir . --reload --port 8001")
add_p(doc, "uvicorn web.keras_app:app --app-dir . --reload --port 8002")
add_p(doc, "Mở http://127.0.0.1:8001 hoặc 8002")

doc.add_page_break()

# CHƯƠNG 9
add_h1(doc, "9. HẠN CHẾ VÀ MỞ RỘNG")
add_h2(doc, "9.1. Hạn chế")
add_p(doc, "• Chỉ một biến (univariate), chưa dùng OHLCV, weekday, holiday.")
add_p(doc, "• Window=30 và epochs=18 cố định, chưa tune.")
add_p(doc, "• Chỉ split 80/20 một lần, chưa walk-forward validation.")
add_p(doc, "• Chưa baseline persistence.")
add_p(doc, "• Kết quả phụ thuộc seed và backend.")

add_h2(doc, "9.2. Hướng phát triển")
add_p(doc, "• Thêm đặc trưng đa biến, chỉ báo kỹ thuật.")
add_p(doc, "• Thử GRU, Bidirectional LSTM, Transformer (PatchTST).")
add_p(doc, "• Walk-forward validation, nhiều seed.")
add_p(doc, "• Docker, CI/CD, authentication, data drift monitoring.")

doc.add_page_break()

# CHƯƠNG 10
add_h1(doc, "10. KẾT LUẬN")
add_h2(doc, "10.1. Đối chiếu yêu cầu")

req_h = ["Yêu cầu", "Trạng thái", "Vị trí"]
req_d = [
    ["Lý thuyết RNN/LSTM, hàm, code", "Hoàn thành", "Chương 2"],
    ["Hai dataset: stock, retail", "Hoàn thành", "Chương 3"],
    ["Mô tả, phân bố dữ liệu", "Hoàn thành", "Chương 3"],
    ["PyTorch + deploy web", "Hoàn thành", "Chương 5, 8"],
    ["Keras + deploy web", "Hoàn thành", "Chương 6, 8"],
    ["So sánh hai mô hình", "Hoàn thành", "Chương 7"]
]
add_table(doc, req_h, req_d)

add_h2(doc, "10.2. Tổng kết")
add_p(doc, "Đã hoàn thành đầy đủ 5 yêu cầu: lý thuyết, hai dataset với EDA, PyTorch + web, Keras + web, so sánh. Pipeline: data -> clean -> split -> scale -> window -> LSTM -> metrics -> API -> web. Kết quả notebook: Keras MAE/RMSE tốt hơn, PyTorch nhanh hơn. Hai API lắng nghe port 8001/8002 với giao diện tương tác.")

doc.add_page_break()

# PHỤ LỤC
add_h1(doc, "PHỤ LỤC A: CÀI ĐẶT VÀ TÁI LẬP")
add_p(doc, "conda activate assignment2")
add_p(doc, "cd /home/dau/assignment2/assignment6")
add_p(doc, "python train_pytorch.py")
add_p(doc, "python train_keras.py")
add_p(doc, "uvicorn web.pytorch_app:app --app-dir . --reload --port 8001")
add_p(doc, "uvicorn web.keras_app:app --app-dir . --reload --port 8002")

add_h1(doc, "PHỤ LỤC B: CẤU TRÚC THƯ MỤC")
add_p(doc, "assignment6/")
add_p(doc, "├── common.py, train_pytorch.py, train_keras.py")
add_p(doc, "├── data/ (all_stocks_5yr.csv.zip, online_retail_II.csv.zip)")
add_p(doc, "├── models/ (.pth, .keras, scaler JSON, comparison CSV)")
add_p(doc, "├── figures/ (9 hình PNG)")
add_p(doc, "├── notebook/ (assignment6_rnn_real_data.ipynb)")
add_p(doc, "└── web/ (pytorch_app.py, keras_app.py, index.html)")

add_h1(doc, "PHỤ LỤC C: TÀI LIỆU THAM KHẢO")
add_p(doc, "[1] Hochreiter & Schmidhuber (1997). Long short-term memory. Neural computation.")
add_p(doc, "[2] Goodfellow et al. (2016). Deep Learning. MIT Press.")
add_p(doc, "[3] Kaggle S&P 500: https://www.kaggle.com/datasets/camnugent/sandp500")
add_p(doc, "[4] Kaggle Online Retail II: https://www.kaggle.com/datasets/mashlyn/online-retail-ii-uci")
add_p(doc, "[5] GitHub: https://github.com/TrVHau/A2_CT_hautv.287.git")

doc.save(str(OUT))
print(f"✅ Báo cáo Word 20+ trang đã được tạo thành công: {OUT}")
